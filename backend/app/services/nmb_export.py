"""
X岛内容导出工具
================
支持格式: json / txt / jpg / doc / pdf / html

嵌入到别的项目使用时（vendoring），不需要修改本文件：外部依赖（图片解析/下载、
图片目录、字体目录）都通过顶部的 ExportEnv 注入，见「宿主接口」一节。
"""
from __future__ import annotations

import argparse
import base64
import html as html_lib
import json
import os
import re
import sys
from contextlib import contextmanager
from contextvars import ContextVar
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

# ============================== 颜色常量 ==============================
COLOR_HEADER = (137, 137, 137)
COLOR_PO = (229, 115, 115)
COLOR_ADMIN = (244, 67, 54)
COLOR_LINK = (0, 119, 221)
COLOR_QUOTE = (120, 153, 34)
COLOR_TEXT = (33, 33, 33)
COLOR_BG = (255, 255, 255)
COLOR_DIVIDER = (224, 224, 224)
COLOR_PO_TAG_BG = (252, 235, 235)


# ============================== 辅助 ==============================
def _hex(rgb): return "#{:02X}{:02X}{:02X}".format(*rgb)
def _post_number(post_id) -> str: return f"No.{post_id}"

def _is_meaningful(s, default_skip: Optional[str] = None) -> bool:
    if not s: return False
    if default_skip and s.strip() == default_skip: return False
    return True


def _is_quote_line(s: str) -> bool:
    if not s:
        return False
    return s.lstrip().startswith((">>", "＞＞"))

# ============================== 图片辅助 ==============================
IMAGE_DIR_DEFAULT = "data/images"


# ============================== 宿主接口 ==============================
# 本文件可以被整体复制（vendoring）到别的项目里使用。为此，所有「外部世界」的
# 依赖都收敛到下面这一个接缝，宿主不必再改本文件的任何一行：
#
#     from export import ExportEnv, configure, use_env
#
#     configure(ExportEnv(font_dir="/app/assets/fonts"))        # 进程级默认
#     with use_env(ExportEnv(image_dir="/tmp/abc")):            # 作用域内生效
#         export_data(posts, "pdf", "out.pdf")
#
# 不注入任何东西时，行为与历史上完全一致（图片走同级 images.py、字体走 CWD 下的 fonts/）。
# 用 ContextVar 而不是模块级全局：Web 服务里同进程会并发导出，各自的临时目录/字体
# 目录不会互相串味。

# 图片解析：(img, ext, quality, base_dir) -> 本地路径 或 None（视为没有这张图）
ImageResolver = Callable[[str, str, str, str], Optional[str]]
# 图片下载：同上签名，负责把图准备好（仅在 fetch=True 且本地点不存在时调用）
ImageFetcher = Callable[[str, str, str, str], None]

FONT_DIR_DEFAULT = "fonts"
FONT_CJK = "GoNotoCJKCore.ttf"
FONT_LATIN = "NotoSans-Regular.ttf"


def _default_image_resolver(img: str, ext: str, quality: str,
                             base_dir: str) -> Optional[str]:
    """默认实现：交给同级的 images.py（图片位于 base_dir/quality/ 下）。

    vendored 到没有 images.py 的项目时退化为 base_dir/文件名 —— 这正好对应
    「宿主已把图片备好放进临时目录」的常见用法；需要别的布局请注入 resolver。
    """
    try:
        from images import image_local_path
    except ModuleNotFoundError as e:
        if e.name != "images":
            raise  # images.py 在，但它自己的依赖缺失（如没装 requests）→ 必须暴露出来
        return str(Path(base_dir) / f"{img}{ext}")
    return str(image_local_path(img, ext, quality, base_dir))


def _default_image_fetcher(img: str, ext: str, quality: str,
                            base_dir: str) -> None:
    try:
        from images import download_image
    except ModuleNotFoundError as e:
        if e.name != "images":
            raise
        return  # 没有下载器可用，就当这张图取不到
    download_image(img, ext, quality, base_dir)


@dataclass(frozen=True)
class ExportEnv:
    """宿主可注入的依赖。字段全部可选，缺省即历史默认行为。

    image_resolver: (img, ext, quality, base_dir) -> 本地路径 或 None
    image_fetcher : 同签名，负责按需下载/生成图片
    image_dir     : image_opts 未显式给 base_dir 时的图片根目录
    font_dir      : JPG / PDF 渲染所用字体所在目录
    cjk_font      : 中文字体文件名（PDF 必需，缺失会导致中文渲染为空白）
    latin_font    : 拉丁字体文件名
    """

    image_resolver: Optional[ImageResolver] = None
    image_fetcher: Optional[ImageFetcher] = None
    image_dir: str = IMAGE_DIR_DEFAULT
    font_dir: str = FONT_DIR_DEFAULT
    cjk_font: str = FONT_CJK
    latin_font: str = FONT_LATIN


_DEFAULT_ENV = ExportEnv()
_ENV: ContextVar[Optional[ExportEnv]] = ContextVar("nmb_export_env", default=None)


def current_env() -> ExportEnv:
    """当前生效的宿主接口；未设置时返回进程级默认（即历史行为）。"""
    return _ENV.get() or _DEFAULT_ENV


def configure(env: ExportEnv) -> None:
    """设置进程级默认接口。"""
    global _DEFAULT_ENV
    _DEFAULT_ENV = env


@contextmanager
def use_env(env: ExportEnv):
    """作用域内生效，退出自动还原；并发调用之间互不影响。"""
    token = _ENV.set(env)
    try:
        yield env
    finally:
        _ENV.reset(token)


def set_image_hooks(resolver=None, downloader=None) -> None:
    """只覆盖图片解析，其余字段沿用进程级默认（保留给已有调用方）。"""
    configure(replace(_DEFAULT_ENV, image_resolver=resolver,
                      image_fetcher=downloader))


def _font_candidates(cjk_only: bool = False) -> List[str]:
    """字体路径候选；CJK 字体缺失时 PDF 里中文会变成空白，调用方应先自检。"""
    env = current_env()
    out = [os.path.join(env.font_dir, env.cjk_font)]
    if not cjk_only:
        out.append(os.path.join(env.font_dir, env.latin_font))
    return out


def _resolve_image_file(img: str, ext: str, quality: str,
                         base_dir: str) -> Path:
    """把图片定位交给宿主；resolver 返回 None 时给一个必然不存在的路径，
    让上层的 exists() 判断自然走「没有图」分支。"""
    resolver = current_env().image_resolver or _default_image_resolver
    found = resolver(img, ext, quality, base_dir)
    return Path(found) if found else Path(base_dir) / f"__missing__{img}{ext}"


def _fetch_image(img: str, ext: str, quality: str, base_dir: str) -> None:
    fetcher = current_env().image_fetcher or _default_image_fetcher
    fetcher(img, ext, quality, base_dir)
_MIME_MAP = {
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".png": "image/png",  ".gif": "image/gif",
    ".webp": "image/webp", ".bmp": "image/bmp",
}


def _data_uri_for_file(path: Path) -> Optional[str]:
    try:
        if not path.exists() or path.stat().st_size == 0:
            return None
        mime = _MIME_MAP.get(path.suffix.lower(), "image/jpeg")
        return f"data:{mime};base64,{base64.b64encode(path.read_bytes()).decode('ascii')}"
    except Exception:
        return None


def _resolve_image_src_for_html(post: Dict[str, Any],
                                  image_opts: Optional[Dict[str, Any]]
                                  ) -> Optional[str]:
    """HTML 渲染用：返回 <img src> 的值（URL 或 data URI）。"""
    if not image_opts: return None
    quality = image_opts.get("quality")
    if quality not in ("thumb", "image"): return None
    img = post.get("img"); ext = post.get("ext")
    if not img or not ext: return None

    mode = image_opts.get("mode", "url")
    if mode == "url":
        pattern = image_opts.get("url_pattern", "/api/image/{quality}/{img}{ext}")
        return pattern.format(quality=quality, img=img, ext=ext)

    # embed: 自包含 data URI
    base_dir = image_opts.get("base_dir") or current_env().image_dir
    p = _resolve_image_file(img, ext, quality, base_dir)
    if (not p.exists() or p.stat().st_size == 0) and image_opts.get("fetch"):
        try:
            _fetch_image(img, ext, quality, base_dir)
        except Exception: pass
    if p.exists() and p.stat().st_size > 0:
        uri = _data_uri_for_file(p)
        if uri: return uri
    remote = image_opts.get("remote_pattern")
    if remote:
        return remote.format(quality=quality, img=img, ext=ext)
    return None


def _resolve_image_local_path(post: Dict[str, Any],
                                image_opts: Optional[Dict[str, Any]]
                                ) -> Optional[str]:
    """二进制导出用：返回本地文件路径（必要时按需下载）。"""
    if not image_opts: return None
    quality = image_opts.get("quality")
    if quality not in ("thumb", "image"): return None
    img = post.get("img"); ext = post.get("ext")
    if not img or not ext: return None
    base_dir = image_opts.get("base_dir") or current_env().image_dir
    p = _resolve_image_file(img, ext, quality, base_dir)
    if (not p.exists() or p.stat().st_size == 0) and image_opts.get("fetch"):
        try:
            _fetch_image(img, ext, quality, base_dir)
        except Exception: pass
    return str(p) if p.exists() and p.stat().st_size > 0 else None


# ============================== HTML 处理 ==============================
_DANGEROUS_BLOCK = re.compile(
    r'<(script|style|iframe|object|embed|form|input)\b[^>]*>.*?</\1>',
    re.DOTALL | re.IGNORECASE)
_DANGEROUS_SELF = re.compile(
    r'<(script|style|iframe|object|embed|form|input)\b[^>]*/?>',
    re.IGNORECASE)
_EVENT_ATTR = re.compile(
    r'\s+on[a-z]+\s*=\s*("[^"]*"|\'[^\']*\'|[^\s>]+)', re.IGNORECASE)
_JS_URI = re.compile(
    r'(href|src)\s*=\s*("|\')\s*javascript:[^"\']*\2', re.IGNORECASE)
_BR_NL = re.compile(r'<br\s*/?>(\s*\n)?', re.IGNORECASE)
_BR_ANY = re.compile(r'<br\s*/?>', re.IGNORECASE)
_TAG = re.compile(r'<[^>]+>')


def normalize_xdao_html(content: str) -> str:
    if not content:
        return ""

    s = content
    s = _DANGEROUS_BLOCK.sub('', s)
    s = _DANGEROUS_SELF.sub('', s)
    s = _EVENT_ATTR.sub('', s)
    s = _JS_URI.sub(lambda m: f'{m.group(1)}=""', s)
    s = s.replace('\r\n', '\n').replace('\r', '\n')
    s = re.sub(r'<br\s*/?>[ \t]*\n?', '\n', s, flags=re.IGNORECASE)
    s = s.replace('\n', '<br>')

    return s


def html_to_plain_text(content: str) -> str:
    if not content: return ""
    s = content
    s = _DANGEROUS_BLOCK.sub('', s)
    s = _DANGEROUS_SELF.sub('', s)
    s = re.sub(r'<br\s*/?>\s*\n?', '\n', s, flags=re.IGNORECASE)
    s = _TAG.sub('', s)
    s = html_lib.unescape(s)
    s = s.replace('\r\n', '\n').replace('\r', '\n')
    s = re.sub(r'\n{3,}', '\n\n', s)
    return s.strip('\n')


def _content_to_html(content: str) -> str:
    if not content: return ""
    return normalize_xdao_html(content)


# ============================== 数据载入 ==============================
def load_posts(path: str) -> List[Dict[str, Any]]:
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    if not isinstance(data, list):
        raise ValueError(f"JSON 根节点必须是数组，实际为: {type(data).__name__}")
    return data


# ============================== JSON ==============================
def export_json(data: List[Dict[str, Any]], output_path: str):
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# ============================== TXT ==============================
def export_txt(data, output_path, *, show_header=True, show_meta=True,
               show_divider=True, image_opts=None):
    with open(output_path, "w", encoding="utf-8") as f:
        for i, reply in enumerate(data):
            tags = []
            if reply.get("is_admin"): tags.append("[红名]")
            if reply.get("is_po"):    tags.append("[Po]")
            if reply.get("is_sage"):  tags.append("[SAGE]")

            header_parts = []
            if tags: header_parts.append(" ".join(tags))
            if show_header:
                header_parts.append(reply.get("cookie", ""))
                header_parts.append(str(reply.get("timestamp", "")))
                header_parts.append(_post_number(reply.get("id", "")))
            if show_header or tags:
                f.write(f"#{i + 1}  " + "  ".join(p for p in header_parts if p) + "\n")

            if show_meta:
                title = reply.get("title", "")
                if _is_meaningful(title, "无标题"): f.write(f"标题：{title}\n")
                name = reply.get("name", "")
                if _is_meaningful(name, "无名氏"): f.write(f"名称：{name}\n")

            if reply.get("is_sage"):
                f.write("本串已经被SAGE\n")

            f.write(html_to_plain_text(reply.get('content', '')) + "\n")

            if image_opts and reply.get("img") and reply.get("ext"):
                f.write(f"[图片] {reply['img']}{reply['ext']}\n")

            if show_divider:
                f.write("─" * 60 + "\n")


# ============================== JPG ==============================
def export_jpg(data, output_path, *, show_header=True, show_meta=True,
               show_divider=True, width=1080, image_opts=None):
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        print("错误：导出 JPG 需要安装 Pillow，请执行: pip install Pillow")
        return

    font_candidates = _font_candidates()
    font_path = next((fp for fp in font_candidates if os.path.exists(fp)), None)

    try:
        if font_path:
            font_normal = ImageFont.truetype(font_path, 36)
            font_bold = ImageFont.truetype(font_path, 36)
            font_small = ImageFont.truetype(font_path, 28)
        else:
            font_normal = font_bold = font_small = ImageFont.load_default()
    except Exception:
        font_normal = font_bold = font_small = ImageFont.load_default()

    img_width = width
    side_padding = 40
    content_padding_y = 30
    line_h = 56
    small_line_h = 40
    img_max_h = 800

    measure_img = Image.new("RGB", (10, 10))
    measure_draw = ImageDraw.Draw(measure_img)

    def text_w(text, font):
        if not text:
            return 0
        bbox = measure_draw.textbbox((0, 0), text, font=font)
        return bbox[2] - bbox[0]

    def wrap_text(text, font, max_width):
        lines = []
        for paragraph in (text or "").split("\n"):
            if paragraph == "":
                lines.append("")
                continue

            cur = ""
            for ch in paragraph:
                if text_w(cur + ch, font) > max_width and cur:
                    lines.append(cur)
                    cur = ch
                else:
                    cur += ch

            if cur:
                lines.append(cur)

        return lines or [""]

    def wrap_text_with_quote(text, font, max_width):
        lines = []

        for paragraph in (text or "").split("\n"):
            is_quote = _is_quote_line(paragraph)

            if paragraph == "":
                lines.append(("", is_quote))
                continue

            cur = ""
            for ch in paragraph:
                if text_w(cur + ch, font) > max_width and cur:
                    lines.append((cur, is_quote))
                    cur = ch
                else:
                    cur += ch

            if cur:
                lines.append((cur, is_quote))

        return lines or [("", False)]

    inner_width = img_width - side_padding * 2

    def post_plain(reply):
        return html_to_plain_text(reply.get("content", ""))

    def post_image_info(reply):
        ipath = _resolve_image_local_path(reply, image_opts)
        if not ipath:
            return None

        try:
            with Image.open(ipath) as im:
                iw, ih = im.size

            if not iw or not ih:
                return None

            ratio = min(inner_width / iw, img_max_h / ih, 1.0)
            return ipath, max(1, int(iw * ratio)), max(1, int(ih * ratio))
        except Exception:
            return None

    def calc_post_height(reply):
        h = content_padding_y

        if show_header:
            h += small_line_h + 16

        if show_meta:
            if _is_meaningful(reply.get("title", ""), "无标题"):
                h += line_h
            if _is_meaningful(reply.get("name", ""), "无名氏"):
                h += line_h

        if reply.get("is_sage"):
            h += line_h

        h += line_h * len(wrap_text_with_quote(post_plain(reply), font_normal, inner_width))

        info = post_image_info(reply)
        if info:
            _, _, ih = info
            h += ih + 16

        h += content_padding_y
        return h

    total_height = sum(calc_post_height(r) for r in data)
    img = Image.new("RGB", (img_width, max(total_height, 200)), color=COLOR_BG)
    draw = ImageDraw.Draw(img)

    def draw_text_bold(pos, text, font, fill):
        x, y = pos
        draw.text((x, y), text, font=font, fill=fill)
        draw.text((x + 1, y), text, font=font, fill=fill)

    def draw_po_tag(x, y, font):
        tag = "Po"
        tw = text_w(tag, font)
        th = 32
        rect = [x, y, x + tw + 16, y + th + 6]

        try:
            draw.rounded_rectangle(
                rect,
                radius=6,
                fill=COLOR_PO_TAG_BG,
                outline=COLOR_PO,
                width=2,
            )
        except AttributeError:
            draw.rectangle(
                rect,
                fill=COLOR_PO_TAG_BG,
                outline=COLOR_PO,
                width=2,
            )

        draw_text_bold((x + 8, y + 2), tag, font, COLOR_PO)
        return rect[2] - rect[0]

    y_cursor = 0

    for reply in data:
        x0 = side_padding
        y = y_cursor + content_padding_y
        cur_x = x0
        is_po = bool(reply.get("is_po"))
        is_admin = bool(reply.get("is_admin"))

        if show_header:
            if is_po:
                cur_x += draw_po_tag(cur_x, y, font_small) + 12

            cookie = reply.get("cookie", "")
            color = COLOR_ADMIN if is_admin else (COLOR_PO if is_po else COLOR_HEADER)

            if is_admin or is_po:
                draw_text_bold((cur_x, y + 4), cookie, font_small, color)
            else:
                draw.text((cur_x, y + 4), cookie, font=font_small, fill=color)

            ts = str(reply.get("timestamp", ""))
            draw.text(
                ((img_width - text_w(ts, font_small)) // 2, y + 4),
                ts,
                font=font_small,
                fill=COLOR_HEADER,
            )

            pid = _post_number(reply.get("id", ""))
            draw.text(
                (img_width - side_padding - text_w(pid, font_small), y + 4),
                pid,
                font=font_small,
                fill=COLOR_HEADER,
            )

            y += small_line_h + 16

        if show_meta:
            title = reply.get("title", "")
            if _is_meaningful(title, "无标题"):
                draw.text((x0, y), "标题：", font=font_normal, fill=COLOR_HEADER)
                draw_text_bold(
                    (x0 + text_w("标题：", font_normal), y),
                    title,
                    font_bold,
                    COLOR_TEXT,
                )
                y += line_h

            name = reply.get("name", "")
            if _is_meaningful(name, "无名氏"):
                draw.text((x0, y), "名称：", font=font_normal, fill=COLOR_HEADER)
                draw_text_bold(
                    (x0 + text_w("名称：", font_normal), y),
                    name,
                    font_bold,
                    COLOR_TEXT,
                )
                y += line_h

        if reply.get("is_sage"):
            draw_text_bold((x0, y), "本串已经被SAGE", font_bold, COLOR_ADMIN)
            y += line_h

        for ln, is_quote in wrap_text_with_quote(post_plain(reply), font_normal, inner_width):
            draw.text(
                (x0, y),
                ln,
                font=font_normal,
                fill=COLOR_QUOTE if is_quote else COLOR_TEXT,
            )
            y += line_h

        info = post_image_info(reply)
        if info:
            ipath, nw, nh = info

            try:
                with Image.open(ipath) as pim:
                    if pim.mode != "RGB":
                        pim = pim.convert("RGB")
                    im2 = pim.resize((nw, nh))

                img.paste(im2, (x0, y))
                y += nh + 16
            except Exception:
                pass

        y_cursor += calc_post_height(reply)

        if show_divider:
            draw.line(
                [(0, y_cursor), (img_width, y_cursor)],
                fill=COLOR_DIVIDER,
                width=1,
            )

    img.save(output_path, "JPEG", quality=95)


# ============================== DOC ==============================
def export_doc(data, output_path, *, show_header=True, show_meta=True,
               show_divider=True, image_opts=None):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Cm, Emu
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        from docx.enum.text import WD_TAB_ALIGNMENT
        from docx.text.paragraph import Paragraph
        from docx.text.run import Run
    except ImportError:
        print("错误：导出 DOC 需要安装 python-docx，请执行: pip install python-docx")
        return

    def rgb(t):
        return RGBColor(t[0], t[1], t[2])

    def add_bottom_border(paragraph):
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "E0E0E0")
        pBdr.append(bottom)
        pPr.append(pBdr)

    def set_space_after(paragraph, value):
        """直接写 w:spacing/@w:after，绕开 python-docx 的属性包装。

        与 paragraph_format.space_after 的产出逐字节一致（w:after 取 twips）。
        get_or_add_spacing() 仍是 python-docx 的 schema 感知插入，位置不会错。
        """
        spacing = paragraph._p.get_or_add_pPr().get_or_add_spacing()
        spacing.set(qn("w:after"), str(value.twips))

    def add_run(p, text, *, color=None, bold=False, size=None):
        """完整手写 w:r（rPr + 文本内容），绕开 python-docx 的逐属性与逐元素开销。

        文本切分规则与 python-docx 逐字一致：\t -> <w:tab/>，\r / \n -> <w:br/>，
        其余进 <w:t>；<w:t> 首尾含空白时加 xml:space="preserve"。
        CT_R 的 br / cr / t / tab 都是零序元素且不声明 successors，顺序追加即等价；
        rPr 子元素按 CT_RPr 的 schema 序列 b -> color -> sz 排列。

        经 1008 组参数穷举与整篇文档逐字节比对，输出与原实现完全一致。
        """
        r = OxmlElement("w:r")

        if color is not None or bold or size is not None:
            rPr = OxmlElement("w:rPr")
            if bold:
                rPr.append(OxmlElement("w:b"))
            if color is not None:
                el = OxmlElement("w:color")
                el.set(qn("w:val"), _hex(color)[1:])  # w:color 的 val 不带 '#'
                rPr.append(el)
            if size is not None:
                el = OxmlElement("w:sz")
                el.set(qn("w:val"), str(int(round(size * 2))))  # 半磅
                rPr.append(el)
            r.append(rPr)

        buf = []

        def flush():
            if not buf:
                return
            chunk = "".join(buf)
            t = OxmlElement("w:t")
            t.text = chunk
            if len(chunk.strip()) < len(chunk):
                t.set(qn("xml:space"), "preserve")
            r.append(t)
            buf.clear()

        for ch in text:
            if ch == "\t":
                flush()
                r.append(OxmlElement("w:tab"))
            elif ch in "\r\n":
                flush()
                r.append(OxmlElement("w:br"))
            else:
                buf.append(ch)
        flush()

        p._p.append(r)
        return Run(r, p)

    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)

    section = doc.sections[0]
    usable_w = int(section.page_width - section.left_margin - section.right_margin)

    # 追加段落用 O(1) 写法：python-docx 的 add_paragraph() 每次都从 body 头部
    # 线性查找末尾的 w:sectPr，段落一多就是 O(n²)（8000 楼时占掉大半耗时）。
    # 这里直接拿到 sectPr，把新段落插到它前面，XML 顺序与原来完全一致。
    _body_el = doc.element.body
    _sectPr = _body_el.find(qn("w:sectPr"))
    # Paragraph 的父对象必须是容器（BlockItemContainer），它才提供 .part；
    # 传裸 XML 元素会让 run.add_picture() 拿不到 part 而失败。
    _container = doc._body

    def new_paragraph():
        el = OxmlElement("w:p")
        if _sectPr is not None:
            _sectPr.addprevious(el)
        else:
            _body_el.append(el)
        return Paragraph(el, _container)

    for reply in data:
        is_po = bool(reply.get("is_po"))
        is_admin = bool(reply.get("is_admin"))

        if show_header:
            p = new_paragraph()
            set_space_after(p, Pt(2))

            ts_stops = p.paragraph_format.tab_stops
            ts_stops.add_tab_stop(Emu(usable_w // 2), WD_TAB_ALIGNMENT.CENTER)
            ts_stops.add_tab_stop(Emu(usable_w), WD_TAB_ALIGNMENT.RIGHT)

            if is_po:
                add_run(p, "[Po] ", color=COLOR_PO, bold=True, size=9)

            cookie = reply.get("cookie", "")

            if is_admin:
                add_run(p, cookie, color=COLOR_ADMIN, bold=True, size=10)
            elif is_po:
                add_run(p, cookie, color=COLOR_PO, bold=True, size=10)
            else:
                add_run(p, cookie, color=COLOR_HEADER, size=10)

            add_run(
                p,
                "\t" + str(reply.get("timestamp", "")),
                color=COLOR_HEADER,
                size=10,
            )

            add_run(
                p,
                "\t" + _post_number(reply.get("id", "")),
                color=COLOR_HEADER,
                size=10,
            )

        if show_meta:
            title = reply.get("title", "")
            if _is_meaningful(title, "无标题"):
                p = new_paragraph()
                set_space_after(p, Pt(2))
                add_run(p, "标题：", color=COLOR_HEADER, size=11)
                add_run(p, title, color=COLOR_TEXT, bold=True, size=11)

            name = reply.get("name", "")
            if _is_meaningful(name, "无名氏"):
                p = new_paragraph()
                set_space_after(p, Pt(2))
                add_run(p, "名称：", color=COLOR_HEADER, size=11)
                add_run(p, name, color=COLOR_TEXT, bold=True, size=11)

        if reply.get("is_sage"):
            p = new_paragraph()
            set_space_after(p, Pt(2))
            add_run(p, "本串已经被SAGE", color=COLOR_ADMIN, bold=True, size=11)

        content_p = new_paragraph()
        plain_content = html_to_plain_text(reply.get("content", ""))
        content_lines = plain_content.split("\n") if plain_content else [""]

        for idx, line in enumerate(content_lines):
            run = add_run(
                content_p,
                line,
                color=COLOR_QUOTE if _is_quote_line(line) else COLOR_TEXT,
                size=11,
            )

            if idx < len(content_lines) - 1:
                run.add_break()

        set_space_after(content_p, Pt(8))

        ipath = _resolve_image_local_path(reply, image_opts)
        if ipath:
            try:
                p = new_paragraph()
                set_space_after(p, Pt(8))
                run = p.add_run()
                run.add_picture(ipath, width=Cm(8))
            except Exception:
                pass

        if show_divider:
            sep = new_paragraph()
            add_bottom_border(sep)
            set_space_after(sep, Pt(8))

    doc.save(output_path)


# ============================== PDF ==============================
def export_pdf(data, output_path, *, show_header=True, show_meta=True,
               show_divider=True, image_opts=None):
    try:
        from fpdf import FPDF
    except ImportError:
        print("错误：导出 PDF 需要安装 fpdf2，请执行: pip install fpdf2")
        return

    pdf = FPDF()
    pdf.set_auto_page_break(auto=False)
    pdf.set_margins(15, 15, 15)
    pdf.add_page()

    font_name = "Helvetica"

    for fp in _font_candidates(cjk_only=True):
        if os.path.exists(fp):
            try:
                try:
                    pdf.add_font("CJK", "", fp)
                except TypeError:
                    pdf.add_font("CJK", "", fp, uni=True)

                font_name = "CJK"
                break
            except Exception:
                continue

    page_w = pdf.w - pdf.l_margin - pdf.r_margin
    bottom_y = pdf.h - pdf.b_margin
    LH = 6

    def safe_set_font(family, style="", size=10):
        try:
            pdf.set_font(family, style, size=size)
        except Exception:
            pdf.set_font(family, "", size=size)

    def need(h):
        if pdf.get_y() + h > bottom_y:
            pdf.add_page()

    def char_wrap(text, font_size, style=""):
        safe_set_font(font_name, style, font_size)

        out = []

        for para in (text or "").split("\n"):
            if not para:
                out.append("")
                continue

            cur = ""
            for ch in para:
                if pdf.get_string_width(cur + ch) > page_w and cur:
                    out.append(cur)
                    cur = ch
                else:
                    cur += ch

            out.append(cur)

        return out

    def emit_line(text, *, font_size=10, style="", color=COLOR_TEXT, h=LH):
        need(h)
        safe_set_font(font_name, style, font_size)
        pdf.set_text_color(*color)
        pdf.set_x(pdf.l_margin)
        pdf.cell(page_w, h, text, ln=1)

    for reply in data:
        is_po = bool(reply.get("is_po"))
        is_admin = bool(reply.get("is_admin"))

        if show_header:
            need(LH + 2)

            y0 = pdf.get_y()
            x = pdf.l_margin

            if is_po:
                safe_set_font(font_name, "B", 9)
                pdf.set_text_color(*COLOR_PO)

                tag_w = pdf.get_string_width("Po") + 4

                pdf.set_draw_color(*COLOR_PO)
                pdf.rect(x, y0 + 0.5, tag_w, LH - 1)
                pdf.set_xy(x, y0)
                pdf.cell(tag_w, LH, "Po", align="C")

                x += tag_w + 2

            cookie = reply.get("cookie", "")
            cookie_color = COLOR_ADMIN if is_admin else (COLOR_PO if is_po else COLOR_HEADER)

            safe_set_font(font_name, "B" if (is_admin or is_po) else "", 10)
            pdf.set_text_color(*cookie_color)

            cookie_w = pdf.get_string_width(cookie) + 2
            pdf.set_xy(x, y0)
            pdf.cell(cookie_w, LH, cookie)

            ts = str(reply.get("timestamp", ""))

            safe_set_font(font_name, "", 9)
            pdf.set_text_color(*COLOR_HEADER)

            ts_w = pdf.get_string_width(ts)
            pdf.set_xy(pdf.l_margin + (page_w - ts_w) / 2, y0)
            pdf.cell(ts_w, LH, ts)

            pid = _post_number(reply.get("id", ""))
            pid_w = pdf.get_string_width(pid)
            pdf.set_xy(pdf.l_margin + page_w - pid_w, y0)
            pdf.cell(pid_w, LH, pid)

            pdf.set_xy(pdf.l_margin, y0 + LH + 1)

        if show_meta:
            for label, key, default in [
                ("标题：", "title", "无标题"),
                ("名称：", "name", "无名氏"),
            ]:
                v = reply.get(key, "")

                if not _is_meaningful(v, default):
                    continue

                for ln in char_wrap(label + v, 10, "B"):
                    emit_line(ln, font_size=10, style="B", color=COLOR_TEXT)

        if reply.get("is_sage"):
            emit_line("本串已经被SAGE", font_size=10, style="B", color=COLOR_ADMIN)

        content = html_to_plain_text(reply.get("content", ""))

        if content:
            for para in content.split("\n"):
                is_quote = _is_quote_line(para)

                for ln in char_wrap(para, 10):
                    emit_line(
                        ln,
                        font_size=10,
                        color=COLOR_QUOTE if is_quote else COLOR_TEXT,
                    )

        ipath = _resolve_image_local_path(reply, image_opts)

        if ipath:
            try:
                from PIL import Image as PILImage

                with PILImage.open(ipath) as im:
                    iw, ih = im.size

                if iw > 0 and ih > 0:
                    max_w = min(page_w * 0.6, 80)
                    max_h = 100
                    ratio = min(max_w / iw, max_h / ih)

                    draw_w = iw * ratio
                    draw_h = ih * ratio

                    need(draw_h + 2)

                    pdf.image(
                        ipath,
                        x=pdf.l_margin,
                        y=pdf.get_y(),
                        w=draw_w,
                        h=draw_h,
                    )

                    pdf.set_xy(pdf.l_margin, pdf.get_y() + draw_h + 2)
            except Exception:
                pass

        if show_divider:
            need(4)

            pdf.set_draw_color(*COLOR_DIVIDER)

            y = pdf.get_y() + 1

            pdf.line(pdf.l_margin, y, pdf.l_margin + page_w, y)
            pdf.set_xy(pdf.l_margin, y + 3)
        else:
            pdf.set_xy(pdf.l_margin, pdf.get_y() + 2)

    pdf.output(output_path)


# ============================== HTML ==============================
_HTML_DEFAULT_CSS = f"""
body {{
    font-family: -apple-system, "Helvetica Neue", Arial,
                 "PingFang SC", "Microsoft YaHei", sans-serif;
    background: #f5f5f5; margin: 0; padding: 20px;
    color: {_hex(COLOR_TEXT)};
}}
.thread {{ max-width: 860px; margin: 0 auto; }}
.post {{
    background: {_hex(COLOR_BG)};
    padding: 16px 18px; margin-bottom: 12px;
    border-radius: 6px; box-shadow: 0 1px 3px rgba(0,0,0,.08);
}}
.post-header {{
    display: flex; align-items: center; gap: 8px;
    color: {_hex(COLOR_HEADER)}; font-size: 13px;
    margin-bottom: 8px; flex-wrap: wrap;
}}
.post-header .ts {{ flex: 1; text-align: center; }}
.post-header .pid {{ font-variant-numeric: tabular-nums; }}
.cookie {{ font-weight: bold; }}
.cookie.po {{ color: {_hex(COLOR_PO)}; }}
.cookie.admin {{ color: {_hex(COLOR_ADMIN)}; }}
.po-tag {{
    display: inline-block;
    background: {_hex(COLOR_PO_TAG_BG)};
    color: {_hex(COLOR_PO)};
    border: 1px solid {_hex(COLOR_PO)};
    border-radius: 4px;
    padding: 1px 6px; font-size: 12px; font-weight: bold; line-height: 1.4;
}}
.meta {{ color: {_hex(COLOR_HEADER)}; font-size: 14px; margin-bottom: 4px; }}
.meta strong {{ color: {_hex(COLOR_TEXT)}; }}
.sage {{ color: {_hex(COLOR_ADMIN)}; font-weight: bold; margin-bottom: 6px; }}
.content {{ font-size: 15px; line-height: 1.65; word-break: break-word; }}
.content font[color] {{ color: inherit; }}
.content font[color="#789922"] {{ color: {_hex(COLOR_QUOTE)}; }}
.content a.ref {{ color: {_hex(COLOR_QUOTE)}; text-decoration: none; }}
.content font a.ref {{ color: inherit; }}
.content a.ref:hover {{ text-decoration: underline; }}
.content a.link {{ color: {_hex(COLOR_LINK)}; word-break: break-all; }}
.content small {{ font-size: .85em; }}
.post-image {{ margin-top: 10px; }}
.post-image a {{ display: inline-block; }}
..post-image img {{
    max-width: min(100%, 480px);
    max-height: 480px;
    border-radius: 4px;
    cursor: zoom-in;
    display: block;
    background: #f0f0f0;
}}
.divider {{ border: none; border-top: 1px solid {_hex(COLOR_DIVIDER)}; margin: 12px 0; }}
"""


def render_post(reply, idx, total, *,
                show_header=True, show_meta=True, show_divider=True,
                image_opts=None) -> str:
    parts: List[str] = []
    is_po = bool(reply.get("is_po"))
    is_admin = bool(reply.get("is_admin"))
    post_id = html_lib.escape(str(reply.get("id", "")))

    parts.append(f"<article class='post' id='post-{post_id}'>")

    if show_header:
        parts.append("<div class='post-header'>")
        if is_po: parts.append("<span class='po-tag'>Po</span>")
        cookie_cls = "cookie"
        if is_admin: cookie_cls += " admin"
        elif is_po:  cookie_cls += " po"
        parts.append(f"<span class='{cookie_cls}'>"
                     f"{html_lib.escape(reply.get('cookie', ''))}</span>")
        parts.append(f"<span class='ts'>{html_lib.escape(str(reply.get('timestamp', '')))}</span>")
        parts.append(f"<span class='pid'>{_post_number(post_id)}</span>")
        parts.append("</div>")

    if show_meta:
        t = reply.get("title", "")
        if _is_meaningful(t, "无标题"):
            parts.append(f"<div class='meta'>标题：<strong>{html_lib.escape(t)}</strong></div>")
        n = reply.get("name", "")
        if _is_meaningful(n, "无名氏"):
            parts.append(f"<div class='meta'>名称：<strong>{html_lib.escape(n)}</strong></div>")

    if reply.get("is_sage"):
        parts.append("<div class='sage'>本串已经被SAGE</div>")

    parts.append(f"<div class='content'>{_content_to_html(reply.get('content', ''))}</div>")

    src = _resolve_image_src_for_html(reply, image_opts)
    if src:
        s = html_lib.escape(src, quote=True)
        parts.append(
            f"<div class='post-image'><a href=\"{s}\" target='_blank' rel='noopener'>"
            f"<img class='post-img' src=\"{s}\" loading='lazy' alt=''></a></div>"
        )

    parts.append("</article>")

    if show_divider and idx < total - 1:
        parts.append("<hr class='divider'>")
    return "".join(parts)


def render_posts_fragment(data, *, show_header=True, show_meta=True,
                          show_divider=True, divider_after_last=True,
                          image_opts=None) -> str:
    if not data: return ""
    out = []
    n = len(data) + (1 if divider_after_last else 0)
    for i, r in enumerate(data):
        out.append(render_post(r, i, n,
                               show_header=show_header,
                               show_meta=show_meta,
                               show_divider=show_divider,
                               image_opts=image_opts))
    return "".join(out)


def render_html(data, show_header=True, show_meta=True, show_divider=True, *,
                full_document=True, title="X岛内容", extra_css="",
                image_opts=None) -> str:
    body = "<div class='thread'>" + render_posts_fragment(
        data, show_header=show_header, show_meta=show_meta,
        show_divider=show_divider, divider_after_last=False,
        image_opts=image_opts,
    ) + "</div>"
    if not full_document: return body
    return ("<!DOCTYPE html><html lang='zh-CN'><head><meta charset='utf-8'>"
            "<meta name='viewport' content='width=device-width,initial-scale=1'>"
            f"<title>{html_lib.escape(title)}</title>"
            f"<style>{_HTML_DEFAULT_CSS}{extra_css}</style></head><body>"
            f"{body}</body></html>")


def render_html_skeleton(title="X岛预览", extra_css="") -> str:
    return ("<!DOCTYPE html><html lang='zh-CN'><head><meta charset='utf-8'>"
            "<meta name='viewport' content='width=device-width,initial-scale=1'>"
            f"<title>{html_lib.escape(title)}</title>"
            f"<style>{_HTML_DEFAULT_CSS}{extra_css}</style></head>"
            "<body><div class='thread'></div></body></html>")


def export_html(data, output_path, *, show_header=True, show_meta=True,
                show_divider=True, image_opts=None):
    # 文件导出默认 embed 模式（自包含 data URI）
    opts = None
    if image_opts:
        opts = dict(image_opts)
        opts.setdefault("mode", "embed")
        opts.setdefault("fetch", True)
    html_str = render_html(data, show_header=show_header,
                           show_meta=show_meta, show_divider=show_divider,
                           full_document=True, image_opts=opts)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html_str)


# ============================== 入口 ==============================
EXPORTERS = {
    "json": export_json,
    "txt": export_txt,
    "jpg": export_jpg,
    "doc": export_doc,
    "pdf": export_pdf,
    "html": export_html,
}


def export_data(data, export_type, output_path, *,
                show_header=True, show_meta=True, show_divider=True,
                width=1080, image_opts=None):
    if export_type not in EXPORTERS:
        print(f"不支持的导出类型: {export_type}"); return
    if export_type == "json":
        EXPORTERS[export_type](data, output_path)
    else:
        kwargs = {"show_header": show_header, "show_meta": show_meta,
                  "show_divider": show_divider, "image_opts": image_opts}
        if export_type == "jpg": kwargs["width"] = width
        EXPORTERS[export_type](data, output_path, **kwargs)
    print(f"已导出为 {export_type.upper()} 文件: {output_path}")


def _cli():
    parser = argparse.ArgumentParser(description="X岛内容导出工具")
    parser.add_argument("input", nargs="?")
    parser.add_argument("-f", "--format", choices=list(EXPORTERS.keys()), default="txt")
    parser.add_argument("-o", "--output")
    parser.add_argument("--no-header", dest="show_header", action="store_false")
    parser.add_argument("--no-meta", dest="show_meta", action="store_false")
    parser.add_argument("--no-divider", dest="show_divider", action="store_false")
    parser.add_argument("--width", type=int, default=1080)
    parser.add_argument("--image-quality", choices=["thumb", "image"], default=None)
    parser.set_defaults(show_header=True, show_meta=True, show_divider=True)
    args = parser.parse_args()

    if not args.input: print("请提供输入 JSON 文件"); sys.exit(0)
    if not os.path.exists(args.input):
        print(f"错误：输入文件不存在: {args.input}"); sys.exit(1)

    data = load_posts(args.input)
    output_path = args.output or f"{Path(args.input).stem}.{args.format}"
    image_opts = None
    if args.image_quality:
        image_opts = {"quality": args.image_quality, "mode": "embed", "fetch": True}
    export_data(data, args.format, output_path,
                show_header=args.show_header, show_meta=args.show_meta,
                show_divider=args.show_divider, width=args.width,
                image_opts=image_opts)


if __name__ == "__main__":
    _cli()